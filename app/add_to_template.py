# To suppress a warning by openpyxl.
import warnings
warnings.simplefilter("ignore")

import sys, os
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.trendline import Trendline, TrendlineLabel
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.text import RichText
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font

def resource_path(relative_path):
    """ Get absolute path to Excel template file. """
    base_path = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base_path, relative_path)

class Template:
    def __init__(self):
        self.collected_messages = ""
        self.solvents = ""
        self.wb = openpyxl.load_workbook(resource_path("HS_Quantification Template (HH v 1.2).xlsx"))

    def create_solvent_sheets(self):
        """ Create a sheet for every solvent in [solvents] """
        from .extract_data_from_files import solvents

        if solvents:
            self.solvents = solvents
        else:
            self.collected_messages = "No CoA of any solvent has been provided!\nExtraction procedure can't continue."
            return(self.collected_messages)
        
        if len(self.solvents) <= 12:
            pass
        else:
            self.collected_messages = "More then 12 solvent CoAs have been provided!\nExtraction procedure can't continue."
            return(self.collected_messages)

        # The specified cells are referenced to sheet "solvent 1" and need to be programmatically added to every solvent sheet.
        cells_with_reference = [
            "I5", "I6", "I7", "I8", "I9", "I10", "I11", "I12", "M21", # Only for sheet "Analytical report"
            "B9", "B10", "B11", "B12", "B13", "B14", "B15", "B16", # 0. Project information
            "A22", "B22", "C22", "D22", # 1. Diluent
            "A38", "B38", "C38", "D38", # 4. Diluted stock solution
            "A43", "A44", "A45", "A46", "A47", "A48", "A49", "A50", "A51", "A52", "A53", "A54", # 5. Blanks
            "A62", "A63", "A64", "A65", "A66", "A67", "A68", "A69", "A70", "A71", "A72", "A73", # 6. Calibration curve - Standard name
            "B62", "B63", "B64", "B65", "B66", "B67", "B68", "B69", "B70", "B71", "B72", "B73", # 6. Calibration curve - V (uL)
            "C62", "C63", "C64", "C65", "C66", "C67", "C68", "C69", "C70", "C71", "C72", "C73", # 6. Calibration curve - From Stock
            "D62", "D63", "D64", "D65", "D66", "D67", "D68", "D69", "D70", "D71", "D72", "D73", # 6. Calibration curve - V added (mL)
            "B78", # Noise
            "A84", "B84", "C84", "D84", # Repeatability and control
            "A95", "A96", "A97", "A98", "A99", # Bracketing control - Control name
            "B95", "B96", "B97", "B98", "B99", # Bracketing control - Bracket
            "C95", "C96", "C97", "C98", "C99", # Bracketing control - V (uL)
            "D95", "D96", "D97", "D98", "D99", # Bracketing control - From Stock
            "E95", "E96", "E97", "E98", "E99", # Bracketing control - V added (mL)
            "A104", "A105", "A111", "A117", "A123", "A129", # 9. Sample(s) - Sample name
            "B105", "B111", "B117", "B123", "B129", # 9. Sample(s) - Bracket
            "C105", "C106", "C107", "C108", "C109", "C110", "C111", "C112", "C113", "C114", "C115", "C116", "C117", "C118", "C119", "C120", "C121", "C122", "C123", "C124", "C125", "C126", "C127", "C128", "C129", "C130", "C131", "C132", "C133", "C134", # 9. Sample(s) - Tag
            "D105", "D106", "D107", "D108", "D109", "D110", "D111", "D112", "D113", "D114", "D115", "D116", "D117", "D118", "D119", "D120", "D121", "D122", "D123", "D124", "D125", "D126", "D127", "D128", "D129", "D130", "D131", "D132", "D133", "D134", # 9. Sample(s) - Weight (mg)
            "E105", "E108", "E109", "E110", "E111", "E114", "E115", "E116", "E117", "E120", "E121", "E122", "E123", "E126", "E127", "E128", "E129", "E132", "E133", "E134", # 9. Sample(s) - V pip (uL)
            "F105", "F108", "F109", "F110", "F111", "F114", "F115", "F116", "F117", "F120", "F121", "F122", "F123", "F126", "F127", "F128", "F129", "F132", "F133", "F134", # 9. Sample(s) - From (Stock)
            "G105", "G106", "G107", "G108", "G109", "G110", "G111", "G112", "G113", "G114", "G115", "G116", "G117", "G118", "G119", "G120", "G121", "G122", "G123", "G124", "G125", "G126", "G127", "G128", "G129", "G130", "G131", "G132", "G133", "G134", # 9. Sample(s) - V diluent (mL)
            "H105", "H106", "H107", "H108", "H109", "H110", "H111", "H112", "H113", "H114", "H115", "H116", "H117", "H118", "H119", "H120", "H121", "H122", "H123", "H124", "H125", "H126", "H127", "H128", "H129", "H130", "H131", "H132", "H133", "H134", # 9. Sample(s) - Sample Conc. (mg/mL)
            "M108", "M109", "M110", "M114", "M115", "M116", "M120", "M121", "M122", "M126", "M127", "M128", "M132", "M133", "M134", # 9. Sample(s) - Nom Spiked Conc (ug/mL)
            ]

        # Add a reference to each cell that needs a reference:
        for j, i in enumerate(self.solvents):
            for z in cells_with_reference[9:]:
                self.wb[f"solvent {j + 2}"][z] = f"='{self.solvents[0]}'!{z}"
            self.wb[f"solvent {j + 1}"].title = i
        
        for z in cells_with_reference[:9]:
                self.wb["Analytical Report"][z] = f"='{self.solvents[0]}'!{z}"
        
        # Remove unused sheets (based on 12 + 1 solvent sheets):
        remaining_sheets = 13 - len(self.solvents)
        for i in range(remaining_sheets):
            self.wb.remove(self.wb[f"solvent {i + len(self.solvents) + 1}"])

        self.solvent_sheets = {i: self.wb[i] for i in self.solvents}

        self.plot_chart()
        
    def plot_chart(self):
        """ Generate a chart and specify the layout for all solvent sheets """

        for i in self.solvents:

            # Find x and y-values:
            xvalues = Reference(self.wb[i], min_col=5, min_row=62, max_row=69)
            yvalues = Reference(self.wb[i], min_col=7, min_row=62, max_row=69)
            
            # Instantiate chart object and set properties:
            chart = ScatterChart()
            chart.x_axis.title = "Concentration (µg/mL)"
            chart.x_axis.majorGridlines = None
            chart.x_axis.numFmt = '0'
            chart.y_axis.title = "Peak area (a.u.*s)"
            chart.y_axis.majorGridlines.spPr = GraphicalProperties(noFill = 'True')
            chart.y_axis.majorGridlines.spPr.ln = LineProperties(solidFill= 'C5E3F5')
            chart.legend = None

            # Instantiate series object and set properties:
            series = Series(yvalues, xvalues)
            series.marker = openpyxl.chart.marker.Marker('circle')
            series.graphicalProperties.line.noFill=True
            line_property = LineProperties(solidFill='A02D96')
            graphical_property = GraphicalProperties(ln=line_property)
            font_property = Font(typeface='Calibri Light')
            character_property = CharacterProperties(latin=font_property, sz=900, solidFill='FF0000')
            paragraph_property = ParagraphProperties(defRPr=character_property)
            richtext_property = RichText(p=[Paragraph(pPr=paragraph_property, endParaRPr=character_property)])
            number_format_property = NumFmt(formatCode='0.0000E+00')
            trendline_label = TrendlineLabel(txPr=richtext_property, numFmt=number_format_property)
            series.trendline = Trendline(dispRSqr=True, dispEq=True, spPr=graphical_property, trendlineLbl=trendline_label)
            chart.series.append(series)

            # Add chart to sheet.
            self.wb[i].add_chart(chart, "N61")
        
        self.add_coa_data()
        
    def add_coa_data(self):
        from .extract_data_from_files import solvents_CoA_data
        
        if "NMP" in solvents_CoA_data.keys():
            self.diluent = "NMP"
        elif "DMAC" in solvents_CoA_data.keys():
            self.diluent = "DMAC"
        elif "DMI" in solvents_CoA_data.keys():
            self.diluent = "DMI"
        else:
            self.diluent = "diluent"
        
        # Add CoA data for 1. Diluent:
        try:
            self.solvent_sheets[self.solvents[0]]["A22"] = self.diluent
            self.solvent_sheets[self.solvents[0]]["B22"] = solvents_CoA_data[self.diluent][0]
            self.solvent_sheets[self.solvents[0]]["C22"] = solvents_CoA_data[self.diluent][1]
            self.solvent_sheets[self.solvents[0]]["D22"] = solvents_CoA_data[self.diluent][2]   
            self.collected_messages += f"The diluent is {self.diluent}!\n\n"
        except:
            self.collected_messages += f"CoA of {self.diluent} not provided or incorect format is used.\n"
        
        # Add CoA for 2. Reference standards:
        for i in self.solvents:

            try:
                self.solvent_sheets[i]["A27"] = i
                self.solvent_sheets[i]["B27"] = solvents_CoA_data[i][0]
                self.solvent_sheets[i]["C27"] = solvents_CoA_data[i][1]
                self.solvent_sheets[i]["D27"] = solvents_CoA_data[i][2]
                self.solvent_sheets[i]["F27"] = solvents_CoA_data[i][3][:3] + " " + solvents_CoA_data[i][3][3:]
                self.solvent_sheets[i]["E27"] = float(solvents_CoA_data[i][4][:-5])
            except:
                self.collected_messages += f"Something went wrong for {i}.\nIs the correct format used for the CoA file?"
                return(self.collected_messages)
        
        self.collected_messages += f"A total of {len(self.solvents)} {'solvents CoAs' if len(self.solvents) > 1 else 'solvent CoA'} has been found.\n"

        self.add_area_height_data_A()

    def add_area_height_data_A(self):

        from .extract_data_from_files import collected_A_files, solvents_area_height_A
        number_of_A_files = len(collected_A_files)        
        
        # 6. Add peak area and peak height data for 6. Calibration curve:
        for i in self.solvents:
        
            if((number_of_A_files) < 8) or (number_of_A_files > 12):
                self.collected_messages += "Between 8 or 12 A-files have to be supplied.\nExtraction procedure can't continue."
                return self.collected_messages
    
            for j in range(number_of_A_files - 4):
                self.solvent_sheets[self.solvents[0]][f"A{73 - j - (12 - number_of_A_files)}"] = f"A{j + 1}"
                self.solvent_sheets[i][f"F{73 - j - (12 - number_of_A_files)}"] = solvents_area_height_A[i][j][0]

            for j in range(4):
                self.solvent_sheets[self.solvents[0]][f"A{65 - j}"] = f"A{j + number_of_A_files - 3}"
                self.solvent_sheets[i][f"F{65 - j}"] = solvents_area_height_A[i][j - 4][0]
                self.solvent_sheets[i][f"I{65 - j}"] = solvents_area_height_A[i][j - 4][1]
            
        self.collected_messages += f"Data of files A1 to A{number_of_A_files} have been transferred succesfully!\n"

        self.add_area_height_data_B()

    def add_area_height_data_B(self):

        from .extract_data_from_files import collected_B_files, solvents_area_height_B
        number_of_B_files = len(collected_B_files)
        
        if not number_of_B_files > 0:
            self.collected_messages += "No B-Files have been supplied!\n"
            self.add_sample_data()
            return(self.collected_messages)
            
        for i in self.solvents:
        
            # Add peak area and retention time data for 7. Repeatability and control:
            for j in range(3):
                self.solvent_sheets[i][f"G{84 + j}"] = solvents_area_height_B[i][j][2]
                self.solvent_sheets[i][f"H{84 + j}"] = solvents_area_height_B[i][j][0]

            # Add peak area for 8. Data for bracketing control:
            for j in range(number_of_B_files - 3):
                self.solvent_sheets[i][f"G{95 + j}"] = solvents_area_height_B[i][3 + j][0]
        
        self.collected_messages += f"Data of files B3.1 to B3.{number_of_B_files} have been transferred succesfully!\n"

        self.add_sample_data()

    def add_sample_data(self):

        from .extract_data_from_files import unique_samples, solvents_area_height_samples
        number_of_unique_samples = len(unique_samples)

        if not number_of_unique_samples > 0:
            self.collected_messages += "No sample data has been supplied.\n"
            self.save_template()
            return(self.collected_messages)
        
        if number_of_unique_samples > 5:
            self.collected_messages += "Between 0 or 5 sample files have to be supplied.\nExtraction procedure can't continue."
            return(self.collected_messages)

        for j in range(len(unique_samples)):
                self.solvent_sheets[self.solvents[0]][f"A{105 + (j * 6)}"] = unique_samples[j]
        
        for i in self.solvents:
            for j in range(len(unique_samples)):

                for z in range(len(solvents_area_height_samples[unique_samples[j]][i])):
                    self.solvent_sheets[i][f"I{105 + z + (j * 6)}"] = solvents_area_height_samples[unique_samples[j]][i][z][0]
                
                # The values for S-A6 and S-A4 have to be manually switched if samples exist with S-files:
                if len(solvents_area_height_samples[unique_samples[j]][i]) == 6:
                    self.solvent_sheets[i][f"I{108 + (j * 6)}"] = solvents_area_height_samples[unique_samples[j]][i][5][0]
                    self.solvent_sheets[i][f"I{110 + (j * 6)}"] = solvents_area_height_samples[unique_samples[j]][i][3][0]

        self.collected_messages += f"Data of {number_of_unique_samples} sample{'s have' if number_of_unique_samples > 1 else ' has'} been transferred succesfully!\n"

        self.save_template()
        
    def save_template(self):
        self.wb.save("HS_Quantification Template (HH v 1.2) (processed).xlsx")
        self.collected_messages += "\nA processed Template file has been created!\n"

    def return_collected_messages(self):
        return self.collected_messages